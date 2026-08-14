import os
import json
import logging
import pandas as pd
from datetime import datetime
from typing import Dict, Any, Optional
from sqlalchemy.orm import Session
from app.db.models import Report, User
from app.services.dataset_analysis_service import DatasetAnalysisService
from app.services.capability_validator import CapabilityValidator
from app.services.grounding_validator import GroundingValidator
from app.services.consistency_validator import ConsistencyValidator

logger = logging.getLogger(__name__)

class MultiFormatReportService:
    @staticmethod
    def generate_report(
        db: Session,
        user_id: int,
        title: str,
        format_type: str = "pdf",
        dataset_id: int = None
    ) -> Report:
        """
        Generates executive reports dynamically from the shared DatasetAnalysisObject.
        """
        os.makedirs("./exports", exist_ok=True)
        file_name = f"report_{user_id}_{format_type.lower()}_{title.replace(' ', '_').lower()}_{int(datetime.utcnow().timestamp())}"
        
        # Load canonical analysis
        if not dataset_id:
            from app.db.models import Dataset
            ds = db.query(Dataset).order_by(Dataset.created_at.desc()).first()
            dataset_id = ds.id if ds else 1

        analysis = DatasetAnalysisService.get_analysis(db, dataset_id)
        domain = analysis["detected_domains"]["primary_domain"]
        row_count = analysis["row_count"]
        quality = analysis["completeness_pct"]
        kpis = analysis["kpis"]
        recs = analysis["recommendations"]

        if format_type.lower() == "excel":
            file_path = f"./exports/{file_name}.xlsx"
            with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                # Dynamic KPIs sheet
                kpi_list = []
                for key, val in kpis.items():
                    kpi_list.append({"Metric": key.replace('_', ' ').title(), "Value": str(val)})
                
                pd.DataFrame(kpi_list).to_excel(writer, sheet_name="Executive Overview", index=False)
                
                # Recommendations sheet
                rec_list = []
                for r in recs:
                    rec_list.append({"Title": r["title"], "Priority": r["priority"], "Strategy Impact": r["impact"]})
                
                pd.DataFrame(rec_list).to_excel(writer, sheet_name="Strategic Recommendations", index=False)

        elif format_type.lower() == "docx":
            file_path = f"./exports/{file_name}.docx"
            try:
                from docx import Document
                doc = Document()
                doc.add_heading(title, 0)
                doc.add_paragraph(f"Report Domain: {domain} | Total Records: {row_count} | Integrity Score: {quality:.1f}%")
                
                doc.add_heading("1. Synthesized KPIs Summary", level=1)
                for key, val in kpis.items():
                    doc.add_paragraph(f"• {key.replace('_',' ').title()}: {val}")
                
                doc.add_heading("2. Dynamic Strategic Recommendations", level=1)
                for r in recs:
                    doc.add_paragraph(f"• [{r['priority'].upper()}] {r['title']}: {r['benefit']}")
                
                doc.save(file_path)
            except Exception:
                # Fallback text
                with open(file_path, "w") as f:
                    f.write(f"Report: {title}\nDomain: {domain}\nKPIs:\n")
                    for k, v in kpis.items():
                        f.write(f"{k}: {v}\n")

        else: # PDF via ReportLab
            file_path = f"./exports/{file_name}.pdf"
            try:
                from reportlab.lib.pagesizes import letter
                from reportlab.lib import colors
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

                doc = SimpleDocTemplate(file_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
                styles = getSampleStyleSheet()
                
                title_style = ParagraphStyle('DT', parent=styles['Heading1'], fontSize=20, leading=24, textColor=colors.HexColor('#1e1b4b'), fontName='Helvetica-Bold')
                subtitle_style = ParagraphStyle('DS', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor('#6366f1'), fontName='Helvetica-Bold', spaceAfter=10)
                heading_style = ParagraphStyle('SH', parent=styles['Heading2'], fontSize=12, leading=16, textColor=colors.HexColor('#312e81'), fontName='Helvetica-Bold', spaceBefore=10, spaceAfter=5)
                body_style = ParagraphStyle('BC', parent=styles['Normal'], fontSize=9, leading=13, textColor=colors.HexColor('#334155'))

                story = [
                    Paragraph(title.upper(), title_style),
                    Paragraph(f"AI Business Intelligence Assistant &mdash; Executive Briefing ({domain})", subtitle_style),
                    HRFlowable(width="100%", thickness=1, color=colors.HexColor('#6366f1'), spaceAfter=10),
                    Paragraph("1. Executive Summary & Core KPIs", heading_style)
                ]

                kpi_data = [[Paragraph("<b>KPI Name</b>", body_style), Paragraph("<b>Grounded Value</b>", body_style)]]
                for key, val in kpis.items():
                    kpi_data.append([
                        Paragraph(key.replace('_', ' ').title(), body_style),
                        Paragraph(f"<b>{val}</b>", body_style)
                    ])
                
                kpi_table = Table(kpi_data, colWidths=[270, 270])
                kpi_table.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#f1f5f9')),
                    ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#e2e8f0')),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
                    ('PADDING', (0,0), (-1,-1), 6),
                ]))
                story.append(kpi_table)
                story.append(Spacer(1, 10))

                story.append(Paragraph("2. Prioritized Strategic Recommendations", heading_style))
                rec_data = [[Paragraph("<b>Priority</b>", body_style), Paragraph("<b>Action Title</b>", body_style), Paragraph("<b>Impact</b>", body_style)]]
                for r in recs:
                    rec_data.append([
                        Paragraph(r["priority"], body_style),
                        Paragraph(r["title"], body_style),
                        Paragraph(r["benefit"], body_style)
                    ])
                rec_table = Table(rec_data, colWidths=[80, 260, 200])
                rec_table.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#f1f5f9')),
                    ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#e2e8f0')),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
                    ('PADDING', (0,0), (-1,-1), 5),
                ]))
                story.append(rec_table)

                doc.build(story)
            except Exception as e:
                with open(file_path, "w") as f:
                    f.write(f"Report: {title}\nDomain: {domain}\nFailed reportlab PDF build.")

        # Save model entry
        report_entry = Report(
            user_id=user_id,
            dataset_id=dataset_id,
            title=title,
            format=format_type.lower(),
            file_path=file_path
        )
        db.add(report_entry)
        db.commit()
        db.refresh(report_entry)

        # Run consistency check
        ConsistencyValidator.validate_consistency("Report", {"kpis": kpis}, analysis)

        return report_entry
